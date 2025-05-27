from langchain_community.llms import Ollama
from langchain.chains import LLMChain
from langchain.chains.combine_documents.stuff import StuffDocumentsChain
from langchain.prompts import PromptTemplate
from langchain.chains import RetrievalQA
from langchain.docstore.document import Document
from langchain_chroma import Chroma
from langchain_huggingface import HuggingFaceEmbeddings
import openpyxl, os
from autoquest.config import config

def load_qa_dict_from_excel(path="data/QA_sheet.xlsx"):
    wb = openpyxl.load_workbook(path)
    sheet = wb.active
    qa_dict = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        question, answer = row
        if question and answer:
            qa_dict[str(question).strip().lower()] = str(answer).strip()
    return qa_dict

def load_agent():
    embedding_model = HuggingFaceEmbeddings(model_name=config["embedding_model"])

    vectordb = Chroma(
        persist_directory=config["chroma_db_dir"],
        embedding_function=embedding_model,
        collection_name=config["chroma_collection"]
    )

    llm = Ollama(model="mistral", temperature=0.3)

    prompt_template = PromptTemplate(
        input_variables=["context", "question"],
        template="""
You are an intelligent AI assistant helping users fill out job application forms.

Your response must strictly follow these rules:

1. If the question provides a list of options (e.g., Yes/No, locations, skill levels), choose **only one option from that list**. Your answer **must match the text exactly**.
2. Never invent new answers. Only choose from what is explicitly given.
3. If no options are provided and the question is about experience, respond with **"3"** or **"4"** only.
4. If it's an open-ended question, return a short, professional response (max 15 words).
5. Never include labels, explanation, or formatting. Just return the raw answer.

Examples:
Q: How many years of experience do you have with Python?
A: 4

Q: Are you comfortable with W2?
Yes
No
A: Yes

Q: What is your proficiency in English?
None
Conversational
Professional
Native
A: Professional

Q: Describe your experience with cloud architecture.
A: 4 years of experience designing scalable AWS infrastructure.

Now, based on the context, answer the following:

Context:
{context}

Question: {question}

Answer:
        """
    )

    llm_chain = LLMChain(llm=llm, prompt=prompt_template)
    combine_chain = StuffDocumentsChain(llm_chain=llm_chain, document_variable_name="context")

    agent = RetrievalQA(
        retriever=vectordb.as_retriever(search_kwargs={"k": 1}),
        combine_documents_chain=combine_chain,
        input_key="question",
        output_key="result",
        return_source_documents=False
    )

    qa_dict = {}
    qa_path = "data/QA_sheet.xlsx"
    if os.path.exists(qa_path):
        qa_dict = load_qa_dict_from_excel(qa_path)
        print("📄 QA_sheet loaded into memory.")
        docs = [Document(page_content=q, metadata={"answer": a}) for q, a in qa_dict.items()]
        vectordb.add_documents(docs)

    return agent, qa_dict, vectordb
